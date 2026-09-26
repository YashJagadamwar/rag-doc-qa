"""
ingestion.py
------------
Loads raw documents (PDF, DOCX, TXT) and splits them into overlapping
chunks that are small enough to embed and retrieve accurately.

Design notes (useful for interview explanations):
- We keep loading and splitting separate from embedding/storage (vectorstore.py)
  so each stage can be tested and swapped independently.
- RecursiveCharacterTextSplitter tries paragraph -> sentence -> word breaks,
  which keeps chunks semantically coherent instead of cutting mid-sentence.
- Every chunk keeps metadata (source filename + chunk index) so answers
  can be traced back to an exact place in an exact document (citations).
"""

import os
from dataclasses import dataclass
from typing import List

from pypdf import PdfReader
import docx
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document


@dataclass
class LoadedFile:
    filename: str
    text: str


def _read_pdf(path: str) -> str:
    reader = PdfReader(path)
    pages = []
    for i, page in enumerate(reader.pages):
        content = page.extract_text() or ""
        pages.append(f"[page {i + 1}]\n{content}")
    return "\n\n".join(pages)


def _read_docx(path: str) -> str:
    d = docx.Document(path)
    return "\n".join(p.text for p in d.paragraphs if p.text.strip())


def _read_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


LOADERS = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".txt": _read_txt,
    ".md": _read_txt,
}


def load_file(path: str) -> LoadedFile:
    ext = os.path.splitext(path)[1].lower()
    if ext not in LOADERS:
        raise ValueError(f"Unsupported file type: {ext}")
    text = LOADERS[ext](path)
    return LoadedFile(filename=os.path.basename(path), text=text)


def load_files(paths: List[str]) -> List[LoadedFile]:
    return [load_file(p) for p in paths]


def chunk_documents(
    files: List[LoadedFile],
    chunk_size: int = 800,
    chunk_overlap: int = 120,
) -> List[Document]:
    """
    Turns loaded files into LangChain Document chunks with source metadata.
    chunk_size/chunk_overlap are exposed as tunables because the "right"
    value is a genuine trade-off:
      - smaller chunks -> more precise retrieval, less context per chunk
      - larger chunks  -> more context, higher chance of irrelevant text
        diluting the embedding
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    all_chunks: List[Document] = []
    for f in files:
        splits = splitter.split_text(f.text)
        for idx, chunk_text in enumerate(splits):
            all_chunks.append(
                Document(
                    page_content=chunk_text,
                    metadata={
                        "source": f.filename,
                        "chunk_id": idx,
                    },
                )
            )
    return all_chunks
